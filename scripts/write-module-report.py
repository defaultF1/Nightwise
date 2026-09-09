"""Build a Word handoff and its Markdown companion from a local JSON record."""
import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]

def build(source):
    data = json.loads(Path(source).read_text(encoding='utf-8'))
    target = ROOT / 'talks' / data['filename']
    doc = Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    sec.left_margin = sec.right_margin = Inches(0.75)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    doc.styles['Title'].font.size = Pt(25)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 2'].font.size = Pt(12)
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.styles['Caption'].font.size = Pt(10)
    doc.styles['Caption'].font.italic = False
    doc.styles['Caption'].font.bold = False
    doc.styles['Subtitle'].font.italic = False
    doc.core_properties.title = data['title']
    doc.core_properties.author = 'Codex for the NightWise project owner'
    doc.add_paragraph(data['title'], 'Title')
    doc.add_paragraph(data['meta'], 'Subtitle')
    lines = ['# '+data['title'], '', data['meta'], '']
    for block in data['blocks']:
        kind = block['type']
        if kind == 'page':
            doc.add_page_break()
            continue
        if kind == 'heading':
            doc.add_heading(block['text'], level=block.get('level',1))
            lines.extend(['## '+block['text'], ''])
        elif kind == 'image_pair':
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            align = OxmlElement('w:textAlignment')
            align.set(qn('w:val'), 'top')
            p._p.get_or_add_pPr().append(align)
            for index, item in enumerate(block['images']):
                if index:
                    p.add_run('    ')
                size = {'height': Inches(block['height'])} if 'height' in block else {'width': Inches(block.get('width', 3.1))}
                p.add_run().add_picture(str(ROOT / item['path']), **size)
                p._p.xpath('.//wp:docPr')[-1].set('descr', item['caption'])
                lines.extend(['!['+item['caption']+'](../'+item['path']+')', ''])
            doc.add_paragraph(block['caption'], 'Caption')
        elif kind == 'image':
            image = ROOT / block['path']
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(image), width=Inches(block.get('width',7)))
            drawing = p._p.xpath('.//wp:docPr')[0]
            drawing.set('descr', block['caption'])
            doc.add_paragraph(block['caption'], 'Caption')
            lines.extend(['!['+block['caption']+'](../'+block['path']+')', ''])
        else:
            doc.add_paragraph(block['text'], 'List Bullet' if kind == 'bullet' else None)
            lines.extend([('- ' if kind == 'bullet' else '') + block['text'], ''])
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('NightWise  |  ')
    run.font.size = Pt(9)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    target.with_suffix('.md').write_text('\n'.join(lines), encoding='utf-8')
    print(target)

if __name__ == '__main__':
    build(sys.argv[1])
